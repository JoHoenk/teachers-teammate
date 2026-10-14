"""Unit tests for teachers_teammate.application.service."""
# pylint: disable=W0613  # unused-argument — pytest injects fixtures by parameter name; not all are used in every test

from __future__ import annotations

import dataclasses
import os
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from teachers_teammate.application import service as service_module
from teachers_teammate.application.commands import ApplicationCommands
from teachers_teammate.application.service import ProcessingApplicationService
from teachers_teammate.infrastructure.state_repository import DocumentState


# ── Helpers ────────────────────────────────────────────────────────────────


def _make_svc(
    *,
    pipeline_factory=None,
    state_repository_factory=None,
    discovery=None,
) -> ProcessingApplicationService:
    mock_pipeline_factory = pipeline_factory or MagicMock(return_value=MagicMock(run=lambda: 0))
    mock_state_factory = state_repository_factory or MagicMock(return_value=MagicMock())
    mock_discovery = discovery or MagicMock()
    return ProcessingApplicationService(
        discovery=mock_discovery,
        pipeline_factory=mock_pipeline_factory,
        state_repository_factory=mock_state_factory,
    )


def _minimal_state(tmp_path: Path, source: Path) -> DocumentState:
    return DocumentState(
        schema_version=2,
        source_path=str(source.resolve()),
        source_hash="abc123",
    )


# ── run_selected / run_preview_only ───────────────────────────────────────


@pytest.mark.unit
def test_run_selected_delegates_to_pipeline_factory(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService with an injected pipeline_factory
    When   run_selected() is called
    Then   pipeline_factory is called once and its run() return value is forwarded
    """
    from tests.conftest import make_config  # noqa: PLC0415

    cfg = make_config(tmp_path)
    mock_pipeline = MagicMock()
    mock_pipeline.run.return_value = 42
    factory = MagicMock(return_value=mock_pipeline)
    svc = _make_svc(pipeline_factory=factory)

    result = svc.run_selected(cfg)

    factory.assert_called_once()
    mock_pipeline.run.assert_called_once()
    assert result == 42


@pytest.mark.unit
def test_run_selected_forwards_progress_callbacks_to_pipeline_factory(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService with an injected pipeline_factory
    When   run_selected() is called with a ProgressCallbacks bundle
    Then   the same bundle is passed to the pipeline factory as ``progress``
    """
    from tests.conftest import make_config  # noqa: PLC0415

    factory = MagicMock(return_value=MagicMock(run=lambda: 0))
    svc = _make_svc(pipeline_factory=factory)
    progress = service_module.ProgressCallbacks(on_stage_started=MagicMock())

    svc.run_selected(make_config(tmp_path), progress=progress)

    assert factory.call_args.kwargs["progress"] is progress


@pytest.mark.unit
def test_run_preview_only_delegates_to_pipeline_factory(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService with an injected pipeline_factory
    When   run_preview_only() is called
    Then   pipeline_factory is called and run_preview_only() return value is forwarded
    """
    from tests.conftest import make_config  # noqa: PLC0415

    cfg = make_config(tmp_path)
    mock_pipeline = MagicMock()
    mock_pipeline.run_preview_only.return_value = 0
    factory = MagicMock(return_value=mock_pipeline)
    svc = _make_svc(pipeline_factory=factory)

    result = svc.run_preview_only(cfg)

    factory.assert_called_once_with(cfg)
    mock_pipeline.run_preview_only.assert_called_once()
    assert result == 0


# ── preprocess_preview ─────────────────────────────────────────────────────


@pytest.mark.unit
@pytest.mark.use_case("Preview_Preprocessing")
def test_preprocess_preview_uses_injected_component_factory(
    tmp_path: Path, sample_png: Path
) -> None:
    """
    Given  a ProcessingApplicationService constructed with a custom component_factory
    When   preprocess_preview() is called
    Then   the factory's build_preprocessor is used instead of the module default

    Regression guard: preprocess_preview() used to call
    stage_builder.default_build_preprocessor directly, so the preview silently
    diverged from a pipeline configured with a non-default factory.
    """
    from teachers_teammate.config import OcrConfig  # noqa: PLC0415
    from teachers_teammate.infrastructure.stage_builder import (  # noqa: PLC0415
        PipelineComponentFactory,
    )

    class _StubPreprocessor:
        def preprocess(self, image_path):
            return image_path, ["stub_step"]

    build_preprocessor = MagicMock(return_value=_StubPreprocessor())
    factory = PipelineComponentFactory(build_preprocessor=build_preprocessor)
    svc = ProcessingApplicationService(
        pipeline_factory=MagicMock(),
        state_repository_factory=MagicMock(),
        component_factory=factory,
    )
    ocr = OcrConfig(preprocess_method="grayscale")

    _original, _processed, steps = svc.preprocess_preview(sample_png, ocr, tmp_path)

    assert steps == ["stub_step"]
    build_preprocessor.assert_called_once_with(tmp_path, False, ocr)


# ── check_connection — engine checks ──────────────────────────────────────


@pytest.mark.unit
@pytest.mark.use_case("Service_Availability_Check")
def test_check_connection_ollama_engine_returns_true_when_models_present(monkeypatch) -> None:
    """
    Given  the ollama engine and OllamaClient.check_connection reports one model available
    When   check_connection(engine='ollama') is called
    Then   returns (connected=True, model_ok=True, <success message>)
    """
    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    monkeypatch.setattr(
        ollama_utils.OllamaClient,
        "check_connection",
        lambda self, model="": (True, True, "✓ Connected — 1 model(s) available"),
    )
    svc = _make_svc()
    ok, _, msg = svc.check_connection(engine="ollama")
    assert ok is True
    assert "1" in msg


@pytest.mark.unit
def test_check_connection_ollama_engine_returns_false_when_no_models() -> None:
    """
    Given  Ollama is reachable but has no models pulled
    When   check_connection(engine='ollama') is called
    Then   returns (connected=True, model_ok=False, <message mentioning Ollama>)
    """
    svc = _make_svc()
    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    with patch.object(
        ollama_utils.OllamaClient,
        "check_connection",
        return_value=(True, False, "✗ Connected to Ollama but no models are pulled"),
    ):
        connected, model_ok, msg = svc.check_connection(engine="ollama")
    assert connected is True
    assert model_ok is False
    assert "Ollama" in msg or "ollama" in msg.lower()


@pytest.mark.unit
@pytest.mark.use_case("Service_Availability_Check")
def test_check_connection_tesseract_returns_true_when_binary_found(monkeypatch) -> None:
    """
    Given  shutil.which returns a non-None path for 'tesseract'
    When   check_connection(engine='tesseract') is called
    Then   returns (connected=True, model_ok=True, <success message>)
    """
    import shutil  # noqa: PLC0415

    monkeypatch.setattr(shutil, "which", lambda _name: "/usr/bin/tesseract")
    svc = _make_svc()
    ok, _, msg = svc.check_connection(engine="tesseract")
    assert ok is True
    assert "tesseract" in msg.lower()


@pytest.mark.unit
def test_check_connection_tesseract_returns_false_when_binary_missing(monkeypatch) -> None:
    """
    Given  shutil.which returns None for 'tesseract'
    When   check_connection(engine='tesseract') is called
    Then   returns (connected=False, model_ok=False, <error message>)
    """
    import shutil  # noqa: PLC0415

    monkeypatch.setattr(shutil, "which", lambda _name: None)
    svc = _make_svc()
    ok, _, msg = svc.check_connection(engine="tesseract")
    assert ok is False
    assert "tesseract" in msg.lower()


@pytest.mark.unit
def test_check_connection_paddleocr_returns_true_when_installed(monkeypatch) -> None:
    """
    Given  importlib.util.find_spec returns a non-None spec for 'paddleocr'
    When   check_connection(engine='paddleocr') is called
    Then   returns (connected=True, model_ok=True, <success message>)
    """
    import importlib.util  # noqa: PLC0415

    monkeypatch.setattr(importlib.util, "find_spec", lambda _name: object())
    svc = _make_svc()
    ok, _, msg = svc.check_connection(engine="paddleocr")
    assert ok is True
    assert "paddleocr" in msg.lower()


@pytest.mark.unit
def test_check_connection_paddleocr_returns_false_when_not_installed(monkeypatch) -> None:
    """
    Given  importlib.util.find_spec returns None for 'paddleocr'
    When   check_connection(engine='paddleocr') is called
    Then   returns (connected=False, model_ok=False, <error message>)
    """
    import importlib.util  # noqa: PLC0415

    monkeypatch.setattr(importlib.util, "find_spec", lambda _name: None)
    svc = _make_svc()
    ok, _, msg = svc.check_connection(engine="paddleocr")
    assert ok is False
    assert "paddleocr" in msg.lower()


# ── check_connection — provider checks ────────────────────────────────────


@pytest.mark.unit
def test_check_connection_provider_ollama_success(monkeypatch) -> None:
    """
    Given  provider='ollama' and OllamaClient.check_connection reports success
    When   check_connection(provider='ollama') is called
    Then   returns (connected=True, model_ok=True, <success message>)
    """
    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    monkeypatch.setattr(
        ollama_utils.OllamaClient,
        "check_connection",
        lambda self, model="": (True, True, "✓ Connected — 1 model(s) available"),
    )
    svc = _make_svc()
    ok, _, msg = svc.check_connection(provider="ollama")
    assert ok is True


@pytest.mark.parametrize(
    ("provider", "env_var"),
    [
        ("openai", "OPENAI_API_KEY"),
        ("anthropic", "ANTHROPIC_API_KEY"),
        ("google", "GOOGLE_API_KEY"),
        ("mistral", "MISTRAL_API_KEY"),
        ("cohere", "COHERE_API_KEY"),
    ],
)
@pytest.mark.unit
@pytest.mark.use_case("Service_Availability_Check")
def test_check_connection_provider_returns_true_when_env_var_set(
    provider: str, env_var: str, monkeypatch
) -> None:
    """
    Given  a known LLM provider and its API key env var is set to a non-empty value
    When   check_connection(provider=<name>) is called
    Then   returns (connected=True, model_ok=True, <message containing the env var name>)
    """
    monkeypatch.setenv(env_var, "sk-test")
    svc = _make_svc()
    ok, _, msg = svc.check_connection(provider=provider)
    assert ok is True
    assert env_var in msg


@pytest.mark.parametrize(
    ("provider", "env_var"),
    [
        ("openai", "OPENAI_API_KEY"),
        ("anthropic", "ANTHROPIC_API_KEY"),
    ],
)
@pytest.mark.unit
def test_check_connection_provider_returns_false_when_env_var_missing(
    provider: str, env_var: str, monkeypatch
) -> None:
    """
    Given  a known LLM provider and its API key env var is not set
    When   check_connection(provider=<name>) is called
    Then   returns (connected=False, model_ok=False, <message containing the env var name>)
    """
    monkeypatch.delenv(env_var, raising=False)
    svc = _make_svc()
    ok, _, msg = svc.check_connection(provider=provider)
    assert ok is False
    assert env_var in msg


@pytest.mark.unit
def test_check_connection_no_engine_no_provider_returns_false() -> None:
    """
    Given  neither engine nor provider are supplied
    When   check_connection() is called with no arguments
    Then   returns (connected=False, model_ok=False, <message>)
    """
    svc = _make_svc()
    ok, _, msg = svc.check_connection()
    assert ok is False
    assert len(msg) > 0


# ── list_ocr_engines / list_supported_suffixes / default_preprocess_for_engine


@pytest.mark.unit
def test_list_ocr_engines_returns_known_engines() -> None:
    """
    Given  a ProcessingApplicationService
    When   list_ocr_engines() is called
    Then   the returned list includes 'ollama' and 'tesseract'
    """
    svc = _make_svc()
    engines = svc.list_ocr_engines()
    assert "ollama" in engines
    assert "tesseract" in engines


@pytest.mark.unit
def test_list_supported_suffixes_returns_frozenset_with_images() -> None:
    """
    Given  a ProcessingApplicationService
    When   list_supported_suffixes() is called
    Then   the returned frozenset includes .png, .pdf, and .txt
    """
    svc = _make_svc()
    suffixes = svc.list_supported_suffixes()
    assert ".png" in suffixes
    assert ".pdf" in suffixes
    assert ".txt" in suffixes


@pytest.mark.unit
def test_default_preprocess_for_engine_returns_string() -> None:
    """
    Given  a known OCR engine name
    When   default_preprocess_for_engine() is called
    Then   a non-empty string is returned
    """
    svc = _make_svc()
    result = svc.default_preprocess_for_engine("tesseract")
    assert isinstance(result, str)
    assert len(result) > 0


# ── _invalidate_from_stage (static helper) ────────────────────────────────


@pytest.mark.unit
def test_invalidate_from_stage_static_matches_by_resolved_source_id(tmp_path: Path) -> None:
    """
    Given  a list of files and selected_source_ids matching one file by resolved path
    When   _invalidate_from_stage() is called
    Then   the matched file is invalidated in the store and its source_id is returned
    """
    source = tmp_path / "doc.png"
    source.touch()
    mock_store = MagicMock()
    mock_store.invalidate_from_stage.return_value = MagicMock()

    run_ids, failures, nothing_matched = ApplicationCommands._invalidate_from_stage(
        selected_source_ids=[str(source.resolve())],
        stage="ocr",
        files=[source],
        store=mock_store,
    )

    mock_store.invalidate_from_stage.assert_called_once_with(source, "ocr")
    assert str(source.resolve()) in run_ids
    assert failures == []
    assert nothing_matched is False


@pytest.mark.unit
def test_invalidate_from_stage_static_returns_nothing_matched_when_no_files_provided() -> None:
    """
    Given  an empty files list
    When   _invalidate_from_stage() is called
    Then   empty lists are returned and nothing_matched is True
    """
    run_ids, failures, nothing_matched = ApplicationCommands._invalidate_from_stage(
        selected_source_ids=["any"],
        stage="correction",
        files=[],
        store=MagicMock(),
    )
    assert run_ids == []
    assert failures == []
    assert nothing_matched is True


@pytest.mark.unit
def test_invalidate_from_stage_static_accumulates_oserror_failures(tmp_path: Path) -> None:
    """
    Given  a store whose invalidate_from_stage raises OSError
    When   _invalidate_from_stage() is called
    Then   the failure is recorded in the failures list and run_ids still contains the source
    """
    source = tmp_path / "bad.png"
    source.touch()
    mock_store = MagicMock()
    mock_store.invalidate_from_stage.side_effect = OSError("disk full")

    run_ids, failures, nothing_matched = ApplicationCommands._invalidate_from_stage(
        selected_source_ids=[str(source.resolve())],
        stage="ocr",
        files=[source],
        store=mock_store,
    )
    assert "bad.png" in failures[0]
    assert str(source.resolve()) in run_ids


# ── record_manual_ocr_edit ────────────────────────────────────────────────


@pytest.mark.unit
def test_record_manual_ocr_edit_delegates_to_store(tmp_path: Path) -> None:
    """
    Given  a valid source file and a store that returns a state
    When   record_manual_ocr_edit() is called
    Then   the store receives the edited text and the updated state is returned
    """
    from tests.conftest import make_config  # noqa: PLC0415

    source = tmp_path / "page.png"
    source.write_text("x")
    state = _minimal_state(tmp_path, source)
    updated = MagicMock()
    mock_store = MagicMock()
    mock_store.load_or_create.return_value = state
    mock_store.record_manual_ocr_edit.return_value = updated

    with patch(
        "teachers_teammate.application.commands.compute_file_hash",
        return_value="abc123",
    ):
        svc = ProcessingApplicationService(
            state_repository_factory=lambda _: mock_store,
        )
        result = svc.record_manual_ocr_edit(
            make_config(tmp_path),
            source=source,
            edited_text="hello world",
            preview_img="/tmp/preview.png",
        )

    mock_store.record_manual_ocr_edit.assert_called_once()
    assert result is updated


@pytest.mark.unit
def test_record_manual_ocr_edit_returns_none_on_oserror(tmp_path: Path) -> None:
    """
    Given  compute_file_hash raises OSError (file unreadable)
    When   record_manual_ocr_edit() is called
    Then   None is returned without raising
    """
    from tests.conftest import make_config  # noqa: PLC0415

    source = tmp_path / "missing.png"

    with patch(
        "teachers_teammate.application.commands.compute_file_hash",
        side_effect=OSError("no such file"),
    ):
        svc = ProcessingApplicationService(
            state_repository_factory=lambda _: MagicMock(),
        )
        result = svc.record_manual_ocr_edit(
            make_config(tmp_path),
            source=source,
            edited_text="text",
            preview_img="",
        )

    assert result is None


# ── record_manual_correction_edit ─────────────────────────────────────────


@pytest.mark.unit
def test_record_manual_correction_edit_persists_text(tmp_path: Path) -> None:
    """
    Given  a valid source file and a store holding an existing OCR-done state
    When   record_manual_correction_edit() is called
    Then   the store records the correction and the updated state is returned
    """
    from tests.conftest import make_config  # noqa: PLC0415

    source = tmp_path / "page.png"
    source.write_text("x")
    state = dataclasses.replace(_minimal_state(tmp_path, source), ocr_done=True, raw_text="raw")
    updated = MagicMock()
    mock_store = MagicMock()
    mock_store.load_or_create.return_value = state
    mock_store.record_manual_correction_edit.return_value = updated

    with patch(
        "teachers_teammate.application.commands.compute_file_hash",
        return_value="abc123",
    ):
        svc = ProcessingApplicationService(
            state_repository_factory=lambda _: mock_store,
        )
        result = svc.record_manual_correction_edit(
            make_config(tmp_path),
            source=source,
            raw_text="raw",
            preview_img="",
            edited_text="corrected",
        )

    mock_store.record_manual_correction_edit.assert_called_once()
    assert result is updated


@pytest.mark.unit
def test_manual_ocr_edit_survives_reconcile(tmp_path: Path) -> None:
    """
    Given  a manual OCR edit recorded for a source file
    When   the cache is reconciled against the same (unchanged) runtime config
    Then   the edited text is preserved rather than flagged stale and wiped
    """
    from tests.conftest import make_config  # noqa: PLC0415

    from teachers_teammate.infrastructure.file_discovery import (  # noqa: PLC0415
        FileDiscovery,
    )
    from teachers_teammate.infrastructure.preview_image_store import (  # noqa: PLC0415
        PreviewImageStore,
    )
    from teachers_teammate.infrastructure.state_repository import (  # noqa: PLC0415
        StateRepository,
    )
    from teachers_teammate.infrastructure.workflow.cache_service import (  # noqa: PLC0415
        CacheReconciliationService,
    )

    source = tmp_path / "note.txt"
    source.write_text("hello")
    cfg = make_config(tmp_path, correction_enabled=False)

    repo = StateRepository(tmp_path / "state")
    commands = ApplicationCommands(discovery=FileDiscovery(), state_store_factory=lambda _c: repo)
    commands.record_manual_ocr_edit(cfg, source=source, edited_text="MANUAL", preview_img="")

    recon = CacheReconciliationService(
        state_repo=repo,
        artifact_store=PreviewImageStore(tmp_path / "art"),
        config=cfg,
    )
    ctx = recon.prepare(source)

    assert ctx.state.ocr_done is True
    assert ctx.state.raw_text == "MANUAL"


# ── _invalidate_from_stage ValueError ────────────────────────────────────


@pytest.mark.unit
def test_invalidate_from_stage_static_invalid_stage_propagates_value_error(tmp_path: Path) -> None:
    """
    Given  a non-empty files list and an invalid stage name
    When   _invalidate_from_stage() is called
    Then   ValueError is raised (normalize_stage rejects the unknown stage name)
    """
    source = tmp_path / "doc.png"
    source.touch()

    with pytest.raises(ValueError, match="Invalid stage"):
        ApplicationCommands._invalidate_from_stage(
            selected_source_ids=[str(source.resolve())],
            stage="nonexistent",
            files=[source],
            store=MagicMock(),
        )


# ── clear_anonymizer_cache ────────────────────────────────────────────────────


@pytest.mark.unit
def test_clear_anonymizer_cache_empties_dict(tmp_path: Path) -> None:
    """
    Given  a service with cached anonymizer entries
    When   clear_anonymizer_cache() is called
    Then   the internal cache is empty
    """
    svc = _make_svc()
    svc._anonymizer_cache[("English", object())] = MagicMock()
    assert len(svc._anonymizer_cache) == 1

    svc.clear_anonymizer_cache()

    assert len(svc._anonymizer_cache) == 0


# ── New service passthroughs ──────────────────────────────────────────────────


@pytest.mark.unit
def test_detect_gpus_delegates_to_infrastructure(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService
    When   detect_gpus() is called
    Then   it delegates to gpu_detector.detect_gpus and returns the result
    """
    from teachers_teammate.infrastructure.gpu_detector import GpuInfo  # noqa: PLC0415

    fake_gpus = [GpuInfo("nvidia", "Test GPU")]
    svc = _make_svc()

    with patch("teachers_teammate.infrastructure.gpu_detector.detect_gpus", return_value=fake_gpus):
        result = svc.detect_gpus()

    assert result == fake_gpus


@pytest.mark.unit
def test_resolve_config_path_delegates(tmp_path: Path) -> None:
    """
    Given  a config file exists at a known path
    When   resolve_config_path() is called with that explicit path as a string
    Then   it returns the resolved Path without error
    """
    cfg = tmp_path / "ocr.toml"
    cfg.write_text("[settings]\n", encoding="utf-8")
    svc = _make_svc()

    result = svc.resolve_config_path(str(cfg))

    assert result == cfg


@pytest.mark.unit
def test_default_config_path_delegates(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService and a patched storage root
    When   default_config_path() is called
    Then   it delegates to storage_root.default_config_path and returns a Path
    """
    expected = tmp_path / "ocr.toml"
    svc = _make_svc()

    with patch(
        "teachers_teammate.application.service._infra_default_config_path",
        return_value=expected,
    ):
        result = svc.default_config_path()

    assert result == expected


@pytest.mark.unit
def test_addon_packages_dir_delegates(tmp_path: Path) -> None:
    """
    Given  a ProcessingApplicationService
    When   addon_packages_dir() is called
    Then   it delegates to addon_manager.get_packages_dir and returns a Path
    """
    svc = _make_svc()

    with patch(
        "teachers_teammate.infrastructure.addon_manager.get_packages_dir",
        return_value=tmp_path,
    ):
        result = svc.addon_packages_dir()

    assert result == tmp_path


@pytest.mark.unit
def test_get_installation_status_reports_tesseract_and_paddleocr() -> None:
    """
    Given  a ProcessingApplicationService and patched shutil.which + importlib.util.find_spec
    When   get_installation_status() is called
    Then   'tesseract' is True when which returns a path and 'paddleocr' is False when spec is None
    """
    svc = _make_svc()

    with (
        patch(
            "teachers_teammate.application.service.shutil.which", return_value="/usr/bin/tesseract"
        ),
        patch("teachers_teammate.application.service.importlib.util.find_spec", return_value=None),
    ):
        status = svc.get_installation_status()

    assert status["tesseract"] is True
    assert status["paddleocr"] is False


@pytest.mark.unit
def test_is_module_importable_returns_true_when_spec_found(monkeypatch) -> None:
    """
    Given  importlib.util.find_spec returns a non-None spec for a module name
    When   is_module_importable() is called with that module name
    Then   True is returned
    """
    import importlib.util  # noqa: PLC0415

    monkeypatch.setattr(importlib.util, "find_spec", lambda _name: object())
    svc = _make_svc()
    assert svc.is_module_importable("some_package") is True


@pytest.mark.unit
def test_is_module_importable_returns_false_when_spec_missing(monkeypatch) -> None:
    """
    Given  importlib.util.find_spec returns None (module not installed)
    When   is_module_importable() is called
    Then   False is returned
    """
    import importlib.util  # noqa: PLC0415

    monkeypatch.setattr(importlib.util, "find_spec", lambda _name: None)
    svc = _make_svc()
    assert svc.is_module_importable("missing_package") is False


@pytest.mark.unit
def test_get_installation_status_includes_langchain_packages(monkeypatch) -> None:
    """
    Given  get_installation_status() is called with a langchain_packages dict
    When   importlib.util.find_spec returns a spec for one module and None for another
    Then   the result contains 'langchain:{key}' entries matching the spec outcomes
    """
    import importlib.util  # noqa: PLC0415

    present_modules = {"langchain_core"}
    monkeypatch.setattr(
        importlib.util,
        "find_spec",
        lambda name: object() if name in present_modules else None,
    )
    svc = _make_svc()
    status = svc.get_installation_status(
        langchain_packages={"core": "langchain_core", "openai": "langchain_openai"},
    )
    assert status["langchain:core"] is True
    assert status["langchain:openai"] is False


@pytest.mark.unit
def test_list_provider_models_uses_ttl_cache() -> None:
    """
    Given  a ProcessingApplicationService whose list_provider_models is called once
    When   list_provider_models() is called a second time within the TTL window
    Then   the underlying infrastructure function is called only once (cache hit)
    """
    svc = _make_svc()
    call_count = 0

    def _fake_list(provider, **_kwargs):
        nonlocal call_count
        call_count += 1
        return ["model-a", "model-b"]

    with patch(
        "teachers_teammate.application.service.list_provider_models", side_effect=_fake_list
    ):
        first = svc.list_provider_models("openai")
        second = svc.list_provider_models("openai")

    assert first == ["model-a", "model-b"]
    assert second == ["model-a", "model-b"]
    assert call_count == 1


@pytest.mark.unit
def test_invalidate_model_cache_forces_refetch() -> None:
    """
    Given  a ProcessingApplicationService with a cached model list for 'openai'
    When   invalidate_model_cache('openai') is called and then list_provider_models again
    Then   the underlying infrastructure function is called a second time
    """
    svc = _make_svc()
    call_count = 0

    def _fake_list(provider, **_kwargs):
        nonlocal call_count
        call_count += 1
        return ["model-a"]

    with patch(
        "teachers_teammate.application.service.list_provider_models", side_effect=_fake_list
    ):
        svc.list_provider_models("openai")
        svc.invalidate_model_cache("openai")
        svc.list_provider_models("openai")

    assert call_count == 2


# ── hardware profiles & model recommendations ───────────────────────────────


@pytest.mark.use_case("Hardware_Profile_Override")
def test_resolve_hardware_profile_manual_setting_skips_detection(monkeypatch) -> None:
    """
    Given  a concrete hardware_profile setting ("16gb")
    When   resolve_hardware_profile is called
    Then   a MANUAL detection is returned and no detection runs
    """
    svc = _make_svc()
    calls = []
    monkeypatch.setattr(
        service_module.ProcessingApplicationService,
        "detect_hardware",
        lambda self, url, **kw: calls.append(url),  # pylint: disable=unnecessary-lambda
    )
    detection = svc.resolve_hardware_profile("16gb", "http://127.0.0.1:11434")
    assert detection.profile == "16gb"
    from teachers_teammate.domain.model_catalog import DetectionReason  # noqa: PLC0415

    assert detection.reason is DetectionReason.MANUAL
    assert calls == []


@pytest.mark.use_case("Model_Recommendation_By_Hardware")
def test_detect_hardware_probes_this_computer_once(monkeypatch) -> None:
    """
    Given  a service with patched GPU detection
    When   detect_hardware is called repeatedly, also with differently written URLs
    Then   the detector runs once (process-lifetime probe); refresh re-runs it
    """
    from teachers_teammate.infrastructure import gpu_detector  # noqa: PLC0415
    from teachers_teammate.domain.model_catalog import DetectionReason, HardwareDetection  # noqa: PLC0415

    calls = []

    def fake_detect_gpus():
        calls.append("gpu")
        return [gpu_detector.GpuInfo("nvidia", "RTX 4060", 8188)]

    monkeypatch.setattr(gpu_detector, "detect_gpus", fake_detect_gpus)
    monkeypatch.setattr(gpu_detector, "detect_apple_unified_memory_gb", lambda: None)

    svc = _make_svc()
    first = svc.detect_hardware("http://127.0.0.1:11434")
    second = svc.detect_hardware("http://127.0.0.1:11434")
    # Editing the URL (trailing slash, host spelling) must not probe again.
    third = svc.detect_hardware("http://localhost:11434/")
    assert first.profile == "8gb"
    assert first.reason is DetectionReason.GPU_DETECTED
    assert second == first
    assert third == first
    assert calls == ["gpu"]

    refreshed = svc.detect_hardware("http://127.0.0.1:11434", refresh=True)
    assert isinstance(refreshed, HardwareDetection)
    assert calls == ["gpu", "gpu"]


@pytest.mark.use_case("Hardware_Profile_Override")
def test_detect_hardware_remote_url_skips_local_gpus(monkeypatch) -> None:
    """
    Given  an Ollama URL pointing at another computer
    When   detect_hardware is called
    Then   the REMOTE profile is returned without consulting local GPUs
    """
    from teachers_teammate.domain.model_catalog import DetectionReason  # noqa: PLC0415

    def exploding_detect_gpus():
        raise AssertionError("local GPU detection must not run for remote URLs")

    from teachers_teammate.infrastructure import gpu_detector  # noqa: PLC0415

    monkeypatch.setattr(gpu_detector, "detect_gpus", exploding_detect_gpus)
    monkeypatch.setattr(gpu_detector, "detect_apple_unified_memory_gb", exploding_detect_gpus)

    svc = _make_svc()
    detection = svc.detect_hardware("http://192.168.1.50:11434")
    assert detection.profile == "8gb"
    assert detection.reason is DetectionReason.REMOTE_SERVER
    assert detection.host == "192.168.1.50"
    assert detection.gpu_name is None


@pytest.mark.use_case("Hardware_Profile_Override")
def test_resolve_hardware_profile_unknown_setting_treated_as_auto(monkeypatch) -> None:
    """
    Given  an invalid hardware_profile setting value
    When   resolve_hardware_profile is called
    Then   the value is treated as "auto" and detection runs
    """
    from teachers_teammate.infrastructure import gpu_detector  # noqa: PLC0415
    from teachers_teammate.domain.model_catalog import DetectionReason  # noqa: PLC0415

    monkeypatch.setattr(
        gpu_detector, "detect_gpus", lambda: [gpu_detector.GpuInfo("nvidia", "RTX 4060", 8188)]
    )
    monkeypatch.setattr(gpu_detector, "detect_apple_unified_memory_gb", lambda: None)

    svc = _make_svc()
    detection = svc.resolve_hardware_profile("bogus", "http://127.0.0.1:11434")
    assert detection.profile == "8gb"
    assert detection.reason is DetectionReason.GPU_DETECTED


@pytest.mark.use_case("Model_Category_Guidance")
def test_classify_ollama_models_uses_catalog_and_caches_capabilities(monkeypatch) -> None:
    """
    Given  a mix of catalogued and custom model tags
    When   classify_ollama_models is called twice
    Then   catalogued tags resolve from the catalog, unknown tags via
           model_capabilities, and the result is cached (no second server call)
    """
    svc = _make_svc()
    caps_calls = []

    class _FakeClient:
        def __init__(self, url: str, *, timeout: float = 0.0) -> None:
            _ = url, timeout

        def model_capabilities(self, tag: str):
            caps_calls.append(tag)
            return frozenset({"vision"}) if "vl" in tag else frozenset({"completion"})

    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    monkeypatch.setattr(ollama_utils, "OllamaClient", _FakeClient)

    tags = ["qwen3-vl:8b", "my-custom-vl:latest", "my-custom-text:latest"]
    result = svc.classify_ollama_models(tags, "http://127.0.0.1:11434")
    assert result["qwen3-vl:8b"] == frozenset({"ocr"})
    # A model that can read images may be used for reading AND writing.
    assert result["my-custom-vl:latest"] == frozenset({"ocr", "general"})
    assert result["my-custom-text:latest"] == frozenset({"general"})
    assert sorted(caps_calls) == ["my-custom-text:latest", "my-custom-vl:latest"]

    # Second call is fully cached — the server is not asked again.
    svc.classify_ollama_models(tags, "http://127.0.0.1:11434")
    assert len(caps_calls) == 2

    # invalidate_model_cache clears the classification cache.
    svc.invalidate_model_cache("ollama", "http://127.0.0.1:11434")
    svc.classify_ollama_models(tags, "http://127.0.0.1:11434")
    assert len(caps_calls) == 4

    # ... also when the URL is written differently (trailing slash).
    svc.on_ollama_models_changed("http://127.0.0.1:11434/")
    svc.classify_ollama_models(tags, "http://127.0.0.1:11434")
    assert len(caps_calls) == 6


@pytest.mark.use_case("Recommended_Models_On_First_Start")
@pytest.mark.parametrize(
    "profile",
    ["cpu", "8gb", "16gb"],
)
def test_first_run_model_settings_return_best_models(profile: str) -> None:
    """
    Given  a hardware profile
    When   first_run_model_settings is called
    Then   ocr_model is the BEST reading model and correction/evaluate are the
           BEST writing model for that profile
    """
    from teachers_teammate.domain.model_catalog import best_model  # noqa: PLC0415

    from typing import cast  # noqa: PLC0415

    from teachers_teammate.domain.model_catalog import HardwareProfile  # noqa: PLC0415

    svc = _make_svc()
    settings = svc.first_run_model_settings(cast(HardwareProfile, profile))
    assert settings["ocr_model"] == best_model("ocr", cast(HardwareProfile, profile))
    assert settings["correction_model"] == best_model("general", cast(HardwareProfile, profile))
    assert settings["evaluate_model"] == best_model("general", cast(HardwareProfile, profile))


@pytest.mark.use_case("Model_Category_Guidance")
def test_unknown_classification_is_not_cached(monkeypatch) -> None:
    """
    Given  a custom model whose capabilities could not be read (None)
    When   it is classified twice
    Then   the server is asked both times — "unknown" is never cached, so a
           model that was not pulled yet is recognised as soon as it is
    """
    svc = _make_svc()
    answers = iter([None, frozenset({"vision"})])
    asked: list[str] = []

    class _FakeClient:
        def __init__(self, url: str, *, timeout: float = 0.0) -> None:
            _ = url, timeout

        def model_capabilities(self, tag: str):
            asked.append(tag)
            return next(answers)

    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    monkeypatch.setattr(ollama_utils, "OllamaClient", _FakeClient)
    url = "http://127.0.0.1:11434"
    assert svc.classify_ollama_models(["my-model:latest"], url) == {"my-model:latest": None}
    assert svc.classify_ollama_models(["my-model:latest"], url) == {
        "my-model:latest": frozenset({"ocr", "general"})
    }
    assert asked == ["my-model:latest", "my-model:latest"]


@pytest.mark.use_case("Model_Category_Guidance")
def test_unreachable_server_is_asked_once_not_per_model(monkeypatch) -> None:
    """
    Given  several custom models and a server that cannot be reached
    When   they are classified
    Then   the server is contacted only for the first model (one timeout, not one per
           model), every model maps to None, and a later call asks again
    """
    from teachers_teammate.exceptions import OllamaConnectionError  # noqa: PLC0415
    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    asked: list[str] = []
    timeouts: list[float] = []

    class _DownClient:
        def __init__(self, url: str, *, timeout: float = 0.0) -> None:
            _ = url
            timeouts.append(timeout)

        def model_capabilities(self, tag: str):
            asked.append(tag)
            raise OllamaConnectionError("down")

    monkeypatch.setattr(ollama_utils, "OllamaClient", _DownClient)
    svc = _make_svc()
    tags = ["a:latest", "b:latest", "c:latest"]
    assert svc.classify_ollama_models(tags, "http://127.0.0.1:11434") == dict.fromkeys(tags)
    assert asked == ["a:latest"]
    assert timeouts and all(0 < t <= 5 for t in timeouts)  # a short, bounded wait

    svc.classify_ollama_models(tags, "http://127.0.0.1:11434")
    assert asked == ["a:latest", "a:latest"]  # nothing was cached


@pytest.mark.use_case("Model_Category_Guidance")
def test_cached_only_classification_never_contacts_the_server(monkeypatch) -> None:
    """
    Given  a custom model that is not classified yet
    When   classify_ollama_models is called with cached_only=True
    Then   the server is not contacted and the model maps to None
    """
    svc = _make_svc()

    def exploding_client(url, **_kw):
        raise AssertionError("cached_only must not contact the server")

    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    monkeypatch.setattr(ollama_utils, "OllamaClient", exploding_client)
    result = svc.classify_ollama_models(
        ["qwen3-vl:8b", "custom:latest"], "http://x:1", cached_only=True
    )
    assert result == {"qwen3-vl:8b": frozenset({"ocr"}), "custom:latest": None}


@pytest.mark.use_case("Model_Recommendation_By_Hardware")
def test_rate_ollama_models_rates_catalog_and_custom_models(monkeypatch) -> None:
    """
    Given  a catalogued model and a custom vision model on the server
    When   rate_ollama_models is called for the reading task on 8 GB
    Then   the catalog model is BEST with its memory/note and the custom one is
           a WORKS model for reading (the single rating path of the widgets)
    """
    from teachers_teammate.domain.model_catalog import Suitability  # noqa: PLC0415
    from teachers_teammate.infrastructure import ollama_utils  # noqa: PLC0415

    class _FakeClient:
        def __init__(self, url: str, *, timeout: float = 0.0) -> None:
            _ = url, timeout

        def model_capabilities(self, tag: str):
            _ = tag
            return frozenset({"vision"})

    monkeypatch.setattr(ollama_utils, "OllamaClient", _FakeClient)
    svc = _make_svc()
    ratings = svc.rate_ollama_models(
        ["qwen3-vl:8b", "my-vision:latest"], "ocr", "8gb", "http://127.0.0.1:11434"
    )
    best, custom = ratings
    assert best.suitability is Suitability.GOOD  # BEST on 8 GB is the default reading model
    assert best.memory_gb == 7.0
    assert best.note
    assert custom.suitability is Suitability.WORKS
    assert custom.model_uses == frozenset({"ocr", "general"})
    assert custom.memory_gb is None


@pytest.mark.use_case("Model_Category_Guidance")
def test_ollama_model_cache_key_ignores_url_spelling(monkeypatch) -> None:
    """
    Given  the model list cached under "http://host:11434/"
    When   it is invalidated through "http://host:11434"
    Then   the entry is gone (both spellings share one cache entry)
    """
    calls: list[str] = []

    def fake_list(provider: str, **_kw) -> list[str]:
        calls.append(provider)
        return ["a:latest"]

    monkeypatch.setattr(service_module, "list_provider_models", fake_list)
    svc = _make_svc()
    svc.list_provider_models("ollama", base_url="http://host:11434/")
    svc.list_provider_models("ollama", base_url="http://host:11434")
    assert len(calls) == 1  # second call was a cache hit
    svc.invalidate_model_cache("ollama", "http://host:11434")
    svc.list_provider_models("ollama", base_url="http://host:11434/")
    assert len(calls) == 2


@pytest.mark.use_case("Model_Category_Guidance")
def test_on_ollama_models_changed_evicts_list_and_classification(monkeypatch) -> None:
    """
    Given  a cached model list and cached classifications
    When   on_ollama_models_changed is called (a pull finished)
    Then   the next lookup asks the server again
    """
    installed = ["a:latest"]
    monkeypatch.setattr(service_module, "list_provider_models", lambda p, **kw: list(installed))
    svc = _make_svc()
    url = "http://127.0.0.1:11434"
    assert svc.list_provider_models("ollama", base_url=url) == ["a:latest"]
    installed.append("b:latest")  # a pull finished
    assert svc.list_provider_models("ollama", base_url=url) == ["a:latest"]  # stale cache
    svc.on_ollama_models_changed(url)
    assert svc.list_provider_models("ollama", base_url=url) == ["a:latest", "b:latest"]


@pytest.mark.use_case("Model_Recommendation_By_Hardware")
def test_installation_status_compares_normalised_tags(monkeypatch) -> None:
    """
    Given  an installed model "llama3.2:latest"
    When   the installation status asks for "llama3.2" (no suffix)
    Then   it counts as installed
    """
    monkeypatch.setattr(service_module, "list_provider_models", lambda p, **kw: ["llama3.2:latest"])
    svc = _make_svc()
    status = svc.get_installation_status(
        ollama_models=["llama3.2", "other:1b"], ollama_url="http://127.0.0.1:11434"
    )
    assert status["ollama:llama3.2"] is True
    assert status["ollama:other:1b"] is False


@pytest.mark.use_case("Model_Recommendation_By_Hardware")
def test_detect_gpus_is_served_from_the_shared_probe(monkeypatch) -> None:
    """
    Given  hardware detection that already probed this computer
    When   detect_gpus() is called afterwards (the "Detect GPU" button)
    Then   the machine is not probed a second time
    """
    from teachers_teammate.infrastructure import gpu_detector  # noqa: PLC0415

    calls: list[str] = []

    def fake_detect_gpus():
        calls.append("gpu")
        return [gpu_detector.GpuInfo("nvidia", "RTX 4060", 8188)]

    monkeypatch.setattr(gpu_detector, "detect_gpus", fake_detect_gpus)
    monkeypatch.setattr(gpu_detector, "detect_apple_unified_memory_gb", lambda: None)
    svc = _make_svc()
    svc.detect_hardware("http://127.0.0.1:11434")
    gpus = svc.detect_gpus()
    assert [g.name for g in gpus] == ["RTX 4060"]
    assert calls == ["gpu"]


@pytest.mark.unit
def test_export_docx_keeps_markdown_formatting(tmp_path: Path) -> None:
    """
    Given  cached Markdown OCR text with a crossed-out word
    When   export_docx() is called with text_format="markdown"
    Then   the DOCX run for that word is struck through and no markers are written
    """
    from docx import Document  # noqa: PLC0415

    svc = _make_svc()
    out = tmp_path / "out" / "doc.docx"

    svc.export_docx(
        tmp_path / "doc.png", "my ~~backpack~~ bag", "", "", out, text_format="markdown"
    )

    cells = Document(str(out)).tables[0].rows[1].cells
    runs = [run for para in cells[1].paragraphs for run in para.runs]
    assert any(run.text == "backpack" and run.font.strike for run in runs)
    assert "~~" not in cells[1].text


@pytest.mark.unit
def test_export_docx_defaults_to_plain_text(tmp_path: Path) -> None:
    """
    Given  cached plain text that contains Markdown-looking characters
    When   export_docx() is called without a text_format
    Then   the characters are written verbatim
    """
    from docx import Document  # noqa: PLC0415

    svc = _make_svc()
    out = tmp_path / "doc.docx"

    svc.export_docx(tmp_path / "doc.png", "~~keep~~", "", "", out)

    assert "~~keep~~" in Document(str(out)).tables[0].rows[1].cells[1].text


@pytest.mark.unit
def test_manual_correction_edit_stores_the_notation_of_a_synthesised_ocr_text(
    tmp_path: Path,
) -> None:
    """
    Given  no state yet and a Markdown OCR text shown in the GUI
    When   the user edits the correction (record_manual_correction_edit)
    Then   the synthesised state stores the OCR text as Markdown, not plain
    """
    from tests.conftest import make_config  # noqa: PLC0415

    from teachers_teammate.infrastructure.file_discovery import (  # noqa: PLC0415
        FileDiscovery,
    )
    from teachers_teammate.infrastructure.state_repository import (  # noqa: PLC0415
        StateRepository,
    )

    source = tmp_path / "note.txt"
    source.write_text("hello")
    repo = StateRepository(tmp_path / "state")
    commands = ApplicationCommands(discovery=FileDiscovery(), state_store_factory=lambda _c: repo)

    state = commands.record_manual_correction_edit(
        make_config(tmp_path),
        source=source,
        raw_text="~~a~~ b",
        preview_img="",
        edited_text="~~a~~ b!",
        raw_text_format="markdown",
    )

    assert state is not None
    assert (state.raw_text, state.raw_text_format) == ("~~a~~ b", "markdown")
    assert state.correction_text == "~~a~~ b!"


# ── Run preflight & model-list caching ──────────────────────────────────────


def _ollama_config(tmp_path: Path, **overrides):
    """Config whose OCR and proofreading both use Ollama."""
    from tests.conftest import make_config  # noqa: PLC0415

    values: dict = {
        "ocr_engine": "ollama",
        "ocr_model": "qwen3-vl:8b",
        "correction_enabled": True,
        "correction_provider": "ollama",
        "correction_model": "qwen3.5:9b",
    }
    values.update(overrides)
    return make_config(tmp_path, **values)


def _server_answers(models: list[str], error: str = ""):
    """Patch the single Ollama request that lists models with ``(models, error)``."""
    return patch(
        "teachers_teammate.infrastructure.ollama_utils.OllamaClient.list_models_or_error",
        return_value=(models, error),
    )


@pytest.mark.unit
def test_empty_model_list_is_not_cached(tmp_path: Path) -> None:
    """
    Given  an Ollama server that is unreachable (empty model list) at first
    When   it is started and the list is requested again
    Then   the second call asks the server again instead of serving the cached failure
    """
    svc = _make_svc()
    answers = [[], ["qwen3-vl:8b"]]

    with patch(
        "teachers_teammate.application.service.list_provider_models",
        side_effect=lambda *_a, **_k: answers.pop(0),
    ):
        first = svc.list_provider_models("ollama", base_url="http://127.0.0.1:11434")
        second = svc.list_provider_models("ollama", base_url="http://127.0.0.1:11434")

    assert first == []
    assert second == ["qwen3-vl:8b"]


@pytest.mark.unit
def test_preflight_reports_unreachable_ollama(tmp_path: Path) -> None:
    """
    Given  a config that needs Ollama and a server that cannot be reached
    When   the run preflight is computed
    Then   it carries the connection message and nothing else
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers([], "✗ Cannot connect to Ollama") as request:
        result = svc.check_run_preflight(cfg)

    request.assert_called_once()  # one timeout, not one per question
    assert result.ok is False
    assert result.ollama_unreachable == "✗ Cannot connect to Ollama"
    assert result.missing_models == ()


@pytest.mark.unit
def test_preflight_lists_missing_models(tmp_path: Path) -> None:
    """
    Given  a reachable server that only has the reading model downloaded
    When   the run preflight is computed
    Then   the proofreading model is reported as missing
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers(["qwen3-vl:8b"]):
        result = svc.check_run_preflight(cfg)

    assert result.missing_models == ("qwen3.5:9b",)
    assert result.ollama_unreachable == ""


@pytest.mark.unit
def test_preflight_connected_server_without_models_reports_all_missing(tmp_path: Path) -> None:
    """
    Given  a reachable Ollama server with no model downloaded at all
    When   the run preflight is computed
    Then   it is not called "unreachable": both configured models are missing
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers([]):
        result = svc.check_run_preflight(cfg)

    assert result.ollama_unreachable == ""
    assert result.missing_models == ("qwen3-vl:8b", "qwen3.5:9b")


@pytest.mark.unit
def test_preflight_ok_when_everything_is_available(tmp_path: Path) -> None:
    """
    Given  a reachable server with all configured models downloaded
    When   the run preflight is computed
    Then   the run may start
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers(["qwen3-vl:8b", "qwen3.5:9b"]):
        result = svc.check_run_preflight(cfg)

    assert result.ok is True


@pytest.mark.unit
def test_preflight_needs_no_ollama_for_other_engines(tmp_path: Path) -> None:
    """
    Given  a config that uses Tesseract and no Ollama stage
    When   the run preflight is computed
    Then   the Ollama server is never contacted and the run may start
    """
    from tests.conftest import make_config  # noqa: PLC0415

    svc = _make_svc()
    cfg = make_config(tmp_path, ocr_engine="tesseract", correction_enabled=False)

    with patch(
        "teachers_teammate.application.service.list_provider_models",
        side_effect=AssertionError("Ollama must not be contacted"),
    ):
        result = svc.check_run_preflight(cfg)

    assert result.ok is True


@pytest.mark.unit
def test_stage_requirements_distinguish_no_models_from_not_running(tmp_path: Path) -> None:
    """
    Given  a server that answers but has no model downloaded
    When   the proofreading requirements are checked
    Then   the issue says no model is downloaded (fixable in Downloads), not "not running"
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers([]):
        issues = svc.check_stage_requirements(cfg, "correction")

    assert [i.message for i in issues] == ["No model is downloaded in Ollama yet."]
    assert issues[0].fixable_via_downloads is True

    with _server_answers([], "✗ Cannot connect to Ollama") as request:
        issues = svc.check_stage_requirements(cfg, "correction")

    request.assert_called_once()
    assert [i.message for i in issues] == ["Ollama is not running. Start it with `ollama serve`."]
    assert issues[0].fixable_via_downloads is False


@pytest.mark.unit
def test_preflight_asks_an_installed_models_server_only_once(tmp_path: Path) -> None:
    """
    Given  a reachable server whose model list was fetched by an earlier check
    When   the run preflight and the stage requirements are computed within the cache TTL
    Then   the server is not asked again
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers(["qwen3-vl:8b", "qwen3.5:9b"]) as request:
        assert svc.check_run_preflight(cfg).ok is True
        assert svc.check_stage_requirements(cfg, "correction") == []

    request.assert_called_once()


@pytest.mark.unit
def test_unreachable_server_is_asked_again_on_the_next_check(tmp_path: Path) -> None:
    """
    Given  a server that was unreachable at the first preflight
    When   it has been started and the preflight is computed again
    Then   the failure was not cached and the run may start
    """
    svc = _make_svc()
    cfg = _ollama_config(tmp_path)

    with _server_answers([], "✗ Cannot connect to Ollama"):
        assert svc.check_run_preflight(cfg).ollama_unreachable
    with _server_answers(["qwen3-vl:8b", "qwen3.5:9b"]):
        assert svc.check_run_preflight(cfg).ok is True
